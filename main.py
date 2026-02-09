from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from typing import Dict, List, Tuple

try:
    import yaml
except Exception:  # pragma: no cover - optional dependency
    yaml = None

from classifier import classify_company
from crawler import fetch_text
from estimator import estimate_purchase
from extractor import extract_info
from translator import translate_to_chinese


def _load_config(config_path: Path) -> Dict:
    if not config_path.exists() or yaml is None:
        return {}
    with config_path.open("r", encoding="utf-8") as handle:
        return yaml.safe_load(handle) or {}


def _read_input(path: Path) -> List[Dict[str, str]]:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        try:
            import pandas as pd

            df = pd.read_excel(path)
            return df.to_dict(orient="records")
        except ImportError:
            from openpyxl import load_workbook

            wb = load_workbook(path)
            sheet = wb.active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            rows = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: row[i] for i in range(len(headers))})
            return rows
    with path.open("r", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        return list(reader)


def _write_excel(path: Path, rows: List[Dict[str, str]]) -> None:
    try:
        import pandas as pd

        df = pd.DataFrame(rows)
        df.to_excel(path, index=False)
        return
    except ImportError:
        from openpyxl import Workbook

        wb = Workbook()
        sheet = wb.active
        if not rows:
            wb.save(path)
            return
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(h) for h in headers])
        wb.save(path)


def _normalize_url(url: str) -> str:
    if not url:
        return ""
    url = url.strip()
    if not url.startswith(("http://", "https://")):
        return f"https://{url}"
    return url


def process_company(row: Dict[str, str], config: Dict) -> Tuple[Dict[str, str], Dict[str, str] | None]:
    company_name = (row.get("company_name") or "").strip()
    website = _normalize_url(row.get("website") or "")
    crawler_cfg = config.get("crawler", {})
    translator_cfg = config.get("translator", {})

    if not website:
        return {}, {"company_name": company_name, "website": website, "reason": "missing_website"}

    crawl = fetch_text(
        website,
        timeout_s=crawler_cfg.get("timeout_s", 20),
        retries=crawler_cfg.get("retries", 2),
        use_playwright_fallback=crawler_cfg.get("use_playwright_fallback", True),
        user_agent=crawler_cfg.get("user_agent", "Mozilla/5.0"),
    )
    if not crawl.ok:
        return {}, {"company_name": company_name, "website": website, "reason": crawl.reason}

    translation = translate_to_chinese(
        crawl.text,
        provider=translator_cfg.get("provider", "libretranslate"),
        endpoint=translator_cfg.get("endpoint"),
        api_key=translator_cfg.get("api_key"),
        timeout_s=translator_cfg.get("timeout_s", 15),
        retries=translator_cfg.get("retries", 1),
    )
    text = translation.text

    extraction = extract_info(text)
    company_type = classify_company(text)

    estimate = estimate_purchase(
        industry=extraction.industry,
        company_type=company_type,
        estimator_cfg=config.get("estimator", {}),
    )

    output = {
        "公司名": company_name,
        "官网": website,
        "年膜类不干胶采购量估计区间": estimate.range_text,
        "估计中位值": estimate.median,
        "行业倾向": extraction.industry,
        "打印机型号/类型": extraction.printer_info,
        "公司类型": company_type,
        "估算说明": estimate.explanation,
        "翻译状态": translation.reason,
    }
    return output, None


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", help="CSV or Excel file with company_name, website columns")
    parser.add_argument("--config", default="config.yaml")
    parser.add_argument("--output", default="output.xlsx")
    parser.add_argument("--manual-review", default="manual_review.xlsx")
    args = parser.parse_args()

    config = _load_config(Path(args.config))
    rows = _read_input(Path(args.input))

    outputs: List[Dict[str, str]] = []
    manual_reviews: List[Dict[str, str]] = []
    for row in rows:
        output, manual = process_company(row, config)
        if manual:
            manual_reviews.append(manual)
        else:
            outputs.append(output)

    _write_excel(Path(args.output), outputs)
    if manual_reviews:
        _write_excel(Path(args.manual_review), manual_reviews)
    return 0


if __name__ == "__main__":
    sys.exit(main())
